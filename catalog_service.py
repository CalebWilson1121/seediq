from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from database import connect, json_dumps, rows_to_dicts
from ingestion import _store_source_document

ALIASES = {
    "product_name": {"product", "product name", "hybrid", "hybrid name", "variety", "variety name", "seed", "seed product"},
    "crop": {"crop", "commodity"},
    "brand": {"brand", "company", "seed brand"},
    "relative_maturity": {"rm", "relative maturity", "maturity", "days"},
    "trait_package": {"trait", "trait package", "traits", "technology", "tech"},
    "yield_ceiling": {"yield", "yield score", "yield potential", "yield ceiling"},
    "drought_score": {"drought", "drought score", "drought tolerance"},
    "wet_soil_score": {"wet soil", "wet soil score", "wet feet", "saturated soils"},
    "emergence_score": {"emergence", "emergence score"},
    "root_score": {"roots", "root", "root score", "root strength"},
    "stalk_score": {"stalk", "stalks", "stalk score", "stalk strength"},
    "disease_score": {"disease", "disease score", "disease package"},
    "placement_text": {"placement", "positioning", "best placement", "management fit", "notes"},
}


def _norm(v: Any) -> str:
    return " ".join(str(v or "").strip().lower().replace("_", " ").split())


def _canonical(header: Any) -> str | None:
    h = _norm(header)
    for key, vals in ALIASES.items():
        if h in vals:
            return key
    return None


def _num(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _tabular_rows(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(x not in (None, "") for x in row)]
    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and isinstance(data.get("products"), list):
            return data["products"]
        if isinstance(data, dict) and isinstance(data.get("rows"), list):
            return data["rows"]
        return [data] if isinstance(data, dict) else []
    raise ValueError(f"Unsupported catalog format: {ext}")


def _mapped_product(raw: dict[str, Any], default_brand: str | None) -> dict[str, Any] | None:
    mapped: dict[str, Any] = {}
    unknown: dict[str, Any] = {}
    for h, value in raw.items():
        key = _canonical(h)
        if key:
            mapped[key] = value
        elif value not in (None, ""):
            unknown[str(h)] = value
    name = str(mapped.get("product_name") or "").strip()
    if not name:
        return None
    return {
        "product_name": name,
        "crop": str(mapped.get("crop") or "").strip() or None,
        "brand": str(mapped.get("brand") or default_brand or "").strip() or None,
        "relative_maturity": _num(mapped.get("relative_maturity")),
        "trait_package": str(mapped.get("trait_package") or "").strip() or None,
        "yield_ceiling": _num(mapped.get("yield_ceiling")),
        "drought_score": _num(mapped.get("drought_score")),
        "wet_soil_score": _num(mapped.get("wet_soil_score")),
        "emergence_score": _num(mapped.get("emergence_score")),
        "root_score": _num(mapped.get("root_score")),
        "stalk_score": _num(mapped.get("stalk_score")),
        "disease_score": _num(mapped.get("disease_score")),
        "placement_text": str(mapped.get("placement_text") or "").strip() or None,
        "metadata": {"unmapped_source_fields": unknown},
    }


def import_catalog(path: Path, original_name: str, organization_id: int, crop_year: int, catalog_name: str, brand: str | None = None) -> dict[str, Any]:
    ext = path.suffix.lower()
    sha = _sha(path)
    stored_path = _store_source_document(path, original_name, sha)
    parser = "seed-catalog-tabular-v0.1"
    status = "draft"
    products: list[dict[str, Any]] = []
    warnings: list[str] = []

    if ext in {".csv", ".xlsx", ".xlsm", ".json"}:
        rows = _tabular_rows(path)
        for i, raw in enumerate(rows, start=2):
            p = _mapped_product(raw, brand)
            if p:
                p["source_row"] = i
                products.append(p)
        if not products:
            warnings.append("No products were recognized. Check the header names or map this catalog manually.")
            status = "needs_mapping"
    elif ext == ".pdf":
        parser = "seed-catalog-pdf-pending-v0.1"
        text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        warnings.append("PDF retained, but product extraction is intentionally waiting for a real dealer catalog layout before we normalize it.")
        status = "needs_mapping"
    else:
        raise ValueError("Seed catalogs currently support CSV, XLSX, XLSM, JSON, or PDF")

    with connect() as conn:
        org = conn.execute("SELECT id,name FROM dealer_organizations WHERE id=?", (organization_id,)).fetchone()
        if not org:
            raise KeyError(f"Dealer organization {organization_id} not found")
        row = conn.execute(
            "INSERT INTO seed_catalogs(organization_id,crop_year,catalog_name,brand,status,source_filename,source_sha256,source_storage_path,parser_name,product_count,metadata_json) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?::jsonb) RETURNING id",
            (organization_id, crop_year, catalog_name, brand, status, original_name, sha, stored_path, parser, len(products), json_dumps({"warnings": warnings})),
        ).fetchone()
        catalog_id = int(row["id"])
        for p in products:
            conn.execute(
                "INSERT INTO seed_products(organization_id,catalog_id,crop_year,crop,brand,company,product_name,relative_maturity,trait_package,yield_ceiling,drought_score,wet_soil_score,emergence_score,root_score,stalk_score,disease_score,placement_text,active,source_row,metadata_json) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (organization_id, catalog_id, crop_year, p["crop"], p["brand"], p["brand"], p["product_name"], p["relative_maturity"], p["trait_package"], p["yield_ceiling"], p["drought_score"], p["wet_soil_score"], p["emergence_score"], p["root_score"], p["stalk_score"], p["disease_score"], p["placement_text"], True, p["source_row"], json_dumps(p["metadata"])),
            )
    return {"catalog_id": catalog_id, "organization_id": organization_id, "crop_year": crop_year, "status": status, "products_imported": len(products), "warnings": warnings, "storage": stored_path, "parser": parser}


def list_organizations() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT id,name,slug,status,created_at FROM dealer_organizations ORDER BY name").fetchall()
    return rows_to_dicts(rows)


def list_catalogs(organization_id: int) -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT * FROM seed_catalogs WHERE organization_id=? ORDER BY crop_year DESC, created_at DESC", (organization_id,)).fetchall()
    return rows_to_dicts(rows)


def list_products(catalog_id: int) -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT * FROM seed_products WHERE catalog_id=? ORDER BY crop, relative_maturity, product_name", (catalog_id,)).fetchall()
    result = rows_to_dicts(rows)
    for row in result:
        if isinstance(row.get("metadata_json"), str):
            try:
                row["metadata_json"] = json.loads(row["metadata_json"])
            except Exception:
                pass
    return result


def publish_catalog(catalog_id: int) -> dict[str, Any]:
    with connect() as conn:
        catalog = conn.execute("SELECT * FROM seed_catalogs WHERE id=?", (catalog_id,)).fetchone()
        if not catalog:
            raise KeyError(f"Catalog {catalog_id} not found")
        count = conn.execute("SELECT COUNT(*) AS n FROM seed_products WHERE catalog_id=?", (catalog_id,)).fetchone()["n"]
        if int(count or 0) == 0:
            raise ValueError("A catalog cannot be published until it has at least one normalized product")
        conn.execute("UPDATE seed_catalogs SET status='published',published_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP,product_count=? WHERE id=?", (int(count), catalog_id))
        return {"catalog_id": catalog_id, "status": "published", "product_count": int(count)}
