from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from models import CropRecord, ParsedDocument, ParsedField, SourceFact

PARSER_VERSION = "0.4.0"

ALIASES = {
    "producer": {"producer", "producer name", "insured", "insured name", "grower"},
    "farm_name": {"farm name", "operation", "operation name"},
    "policy_number": {"policy", "policy number", "policy #"},
    "county": {"county"},
    "state": {"state"},
    "farm_number": {"farm", "farm number", "farm #", "fsa farm"},
    "tract_number": {"tract", "tract number", "tract #"},
    "field_number": {"field", "field number", "field #"},
    "field_name": {"field name", "common name"},
    "crop": {"crop", "commodity"},
    "practice": {"practice", "type/practice"},
    "irrigation": {"irrigation", "irr/non-irr", "irrigated"},
    "acres": {"acres", "reported acres", "planted acres", "insured acres"},
    "crop_year": {"crop year", "year"},
    "production": {"production", "total production"},
    "yield_value": {"yield", "actual yield", "yield/acre"},
    "approved_yield": {"approved yield", "aph", "t-yield", "approved aph"},
    "coverage_level": {"coverage", "coverage level"},
    "unit_structure": {"unit", "unit structure", "unit type"},
}


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _canonical(header: str) -> str | None:
    h = _norm_header(header)
    for key, names in ALIASES.items():
        if h in names:
            return key
    return None


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def detect_document_type(filename: str, preview: str = "") -> str:
    text = f"{filename} {preview[:4000]}".lower()
    if "schedule of insurance" in text and "available units for map view" in text:
        return "SOI"
    if "mbar" in text or "acreage report" in text:
        return "MBAR"
    if "summary of insurance" in text or re.search(r"\bsoi\b", text):
        return "SOI"
    if "aph" in text or "actual production history" in text or "approved yield" in text:
        return "APH"
    return "UNKNOWN"


def read_tabular(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and isinstance(data.get("rows"), list):
            return data["rows"]
        return [data]
    raise ValueError(f"Unsupported tabular extension: {ext}")


def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext in {".txt", ".log", ".csv", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace")
    return ""


def _mapped_row(row: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for header, value in row.items():
        key = _canonical(header)
        if key:
            result[key] = value
    return result


def parse_structured_rows(rows: list[dict[str, Any]], doc_type: str) -> ParsedDocument:
    out = ParsedDocument(document_type=doc_type)
    seen_fields: dict[str, ParsedField] = {}
    for idx, raw in enumerate(rows, start=2):
        row = _mapped_row(raw)
        out.producer_name = out.producer_name or str(row.get("producer") or "").strip() or None
        out.farm_name = out.farm_name or str(row.get("farm_name") or "").strip() or None
        out.policy_number = out.policy_number or str(row.get("policy_number") or "").strip() or None
        fnum = str(row.get("farm_number") or "").strip()
        tnum = str(row.get("tract_number") or "").strip()
        fldnum = str(row.get("field_number") or "").strip()
        fname = str(row.get("field_name") or "").strip()
        field_key = "|".join(x for x in [fnum, tnum, fldnum] if x) or fname or f"row-{idx}"
        field_name = fname or (f"F{fnum} T{tnum} Field {fldnum}" if any([fnum, tnum, fldnum]) else f"Field {idx-1}")
        if field_key not in seen_fields:
            seen_fields[field_key] = ParsedField(name=field_name, acres=_num(row.get("acres")), county=str(row.get("county") or "").strip() or None, state=str(row.get("state") or "").strip() or None, farm_number=fnum or None, tract_number=tnum or None, field_number=fldnum or None, crop=str(row.get("crop") or "").strip() or None, practice=str(row.get("practice") or "").strip() or None, irrigation=str(row.get("irrigation") or "").strip() or None)
        record = CropRecord(field_key=field_key, crop_year=int(_num(row.get("crop_year"))) if _num(row.get("crop_year")) is not None else None, crop=str(row.get("crop") or "").strip() or None, practice=str(row.get("practice") or "").strip() or None, planted_acres=_num(row.get("acres")), production=_num(row.get("production")), yield_value=_num(row.get("yield_value")), approved_yield=_num(row.get("approved_yield")), coverage_level=(_num(row.get("coverage_level")) / 100.0 if (_num(row.get("coverage_level")) or 0) > 1 else _num(row.get("coverage_level"))), unit_structure=str(row.get("unit_structure") or "").strip() or None)
        if any(v is not None for v in [record.crop_year, record.crop, record.planted_acres, record.production, record.yield_value, record.approved_yield, record.coverage_level, record.unit_structure]):
            out.crop_records.append(record)
        for canonical, value in row.items():
            if value not in (None, ""):
                out.facts.append(SourceFact(entity_type="field", entity_key=field_key, field_name=canonical, value=value, source_locator=f"row:{idx}", confidence=1.0))
    out.fields = list(seen_fields.values())
    return out


def parse_loose_text(text: str, doc_type: str) -> ParsedDocument:
    out = ParsedDocument(document_type=doc_type, raw_preview=text[:6000])
    patterns = {"producer_name": r"(?:Producer|Insured)\s*[:#]?\s*([^\n]+)", "policy_number": r"Policy(?: Number| #)?\s*[:#]?\s*([A-Za-z0-9-]+)"}
    for attr, pattern in patterns.items():
        m = re.search(pattern, text, flags=re.I)
        if m:
            setattr(out, attr, m.group(1).strip())
    out.warnings.append("Generic text/PDF parsing only. Add carrier/form-specific parser before production use.")
    return out


def parse_document(path: Path, forced_type: str | None = None) -> tuple[ParsedDocument, str]:
    ext = path.suffix.lower()
    preview = read_text(path)
    doc_type = (forced_type or detect_document_type(path.name, preview)).upper()

    if ext == ".pdf":
        from nau_mapped_soi_parser import looks_like_nau_mapped_soi, parse_nau_mapped_soi_pdf
        if looks_like_nau_mapped_soi(preview):
            return parse_nau_mapped_soi_pdf(path), "nau-mapped-soi-v0.1"

    if doc_type == "MBAR":
        from mbar_parser import parse_mbar_pdf_text, parse_mbar_rows
        if ext in {".csv", ".xlsx", ".xlsm", ".json"}:
            parsed = parse_mbar_rows(read_tabular(path))
            parsed.raw_preview = preview[:6000] if preview else None
            return parsed, "mbar-fields-v0.1"
        if ext == ".pdf":
            return parse_mbar_pdf_text(preview), "mbar-pdf-starter-v0.1"

    if ext in {".csv", ".xlsx", ".xlsm", ".json"}:
        parsed = parse_structured_rows(read_tabular(path), doc_type)
        parsed.raw_preview = preview[:6000] if preview else None
        return parsed, "structured-tabular"
    if ext == ".pdf":
        from nau_aph_parser import looks_like_nau_aph, parse_nau_aph_pdf
        if looks_like_nau_aph(preview):
            return parse_nau_aph_pdf(path), "nau-aph-v0.2"
        return parse_loose_text(preview, doc_type), "generic-text"
    if ext in {".txt", ".log"}:
        return parse_loose_text(preview, doc_type), "generic-text"
    raise ValueError(f"Unsupported file type: {ext}")
