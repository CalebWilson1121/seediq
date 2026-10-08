from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from database import connect, row_to_dict, rows_to_dicts


def _money(value: Any) -> float | None:
    if value is None:
        return None
    return round(float(value), 2)


def _norm_header(value: Any) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def import_price_sheet(path: Path, organization_id: int, crop_year: int) -> dict[str, Any]:
    """Import a dealer price sheet using common spreadsheet headers.

    Required: Product/Product Name and Base Price/Dealer Price.
    Optional: List Price/MSRP and Dealer Cost/Cost.
    """
    rows: list[dict[str, Any]] = []
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            raise ValueError("Pricing spreadsheet is empty")
        headers = [_norm_header(x) for x in values[0]]
        for raw in values[1:]:
            rows.append({headers[i]: raw[i] for i in range(min(len(headers), len(raw)))})
    elif suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for raw in reader:
                rows.append({_norm_header(k): v for k, v in raw.items()})
    else:
        raise ValueError("Upload an .xlsx or .csv pricing file")

    product_keys = ("product", "productname", "hybrid", "variety", "seedproduct", "sku")
    base_keys = ("baseprice", "dealerprice", "standardprice", "farmerbaseprice", "price")
    list_keys = ("listprice", "msrp", "retailprice", "suggestedretail")
    cost_keys = ("dealercost", "cost", "netcost")

    def pick(row: dict[str, Any], keys: tuple[str, ...]):
        for key in keys:
            if key in row and row[key] not in (None, ""):
                return row[key]
        return None

    with connect() as conn:
        products = conn.execute(
            "SELECT sp.id,sp.product_name FROM seed_products sp JOIN seed_catalogs sc ON sc.id=sp.catalog_id "
            "WHERE sc.organization_id=? AND sc.crop_year=?",
            (organization_id, crop_year),
        ).fetchall()
    product_map = {str(p["product_name"]).strip().lower(): int(p["id"]) for p in products}

    imported = 0
    unmatched: list[str] = []
    invalid: list[str] = []
    for row in rows:
        name = str(pick(row, product_keys) or "").strip()
        if not name:
            continue
        product_id = product_map.get(name.lower())
        if not product_id:
            unmatched.append(name)
            continue
        try:
            base_raw = pick(row, base_keys)
            if base_raw in (None, ""):
                invalid.append(name)
                continue
            base_price = float(str(base_raw).replace("$", "").replace(",", ""))
            list_raw = pick(row, list_keys)
            cost_raw = pick(row, cost_keys)
            list_price = float(str(list_raw).replace("$", "").replace(",", "")) if list_raw not in (None, "") else None
            dealer_cost = float(str(cost_raw).replace("$", "").replace(",", "")) if cost_raw not in (None, "") else None
            upsert_dealer_price(organization_id, crop_year, product_id, list_price, base_price, dealer_cost)
            imported += 1
        except Exception:
            invalid.append(name)

    return {
        "crop_year": crop_year,
        "rows_read": len(rows),
        "imported": imported,
        "unmatched_count": len(unmatched),
        "unmatched": unmatched[:50],
        "invalid_count": len(invalid),
        "invalid": invalid[:50],
    }


def list_dealer_prices(organization_id: int, crop_year: int) -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            "SELECT dsp.*,sp.product_name,sp.brand,sp.crop,sp.unit_size_seeds "
            "FROM dealer_seed_prices dsp JOIN seed_products sp ON sp.id=dsp.seed_product_id "
            "WHERE dsp.organization_id=? AND dsp.crop_year=? ORDER BY sp.crop,sp.product_name",
            (organization_id, crop_year),
        ).fetchall()
    return rows_to_dicts(rows)


def upsert_dealer_price(
    organization_id: int,
    crop_year: int,
    seed_product_id: int,
    list_price: float | None,
    base_price: float,
    dealer_cost: float | None,
) -> dict[str, Any]:
    with connect() as conn:
        product = conn.execute(
            "SELECT sp.id,sp.product_name,sp.brand,sp.crop,sc.organization_id,sc.crop_year "
            "FROM seed_products sp JOIN seed_catalogs sc ON sc.id=sp.catalog_id WHERE sp.id=?",
            (seed_product_id,),
        ).fetchone()
        if not product:
            raise KeyError("Seed product not found")
        if int(product["organization_id"]) != int(organization_id):
            raise PermissionError("Seed product belongs to another dealership")
        conn.execute(
            "INSERT INTO dealer_seed_prices(organization_id,crop_year,seed_product_id,list_price,base_price,dealer_cost,status) "
            "VALUES(?,?,?,?,?,?,'active') "
            "ON CONFLICT(organization_id,crop_year,seed_product_id) DO UPDATE SET "
            "list_price=excluded.list_price,base_price=excluded.base_price,dealer_cost=excluded.dealer_cost,status='active',updated_at=CURRENT_TIMESTAMP",
            (organization_id, crop_year, seed_product_id, list_price, base_price, dealer_cost),
        )
        row = conn.execute(
            "SELECT dsp.*,sp.product_name,sp.brand,sp.crop FROM dealer_seed_prices dsp "
            "JOIN seed_products sp ON sp.id=dsp.seed_product_id "
            "WHERE dsp.organization_id=? AND dsp.crop_year=? AND dsp.seed_product_id=?",
            (organization_id, crop_year, seed_product_id),
        ).fetchone()
    return row_to_dict(row) or {}


def get_farmer_profile(farm_id: int, crop_year: int) -> dict[str, Any]:
    with connect() as conn:
        row = conn.execute(
            "SELECT fpp.*,f.farm_name,f.organization_id FROM farms f "
            "LEFT JOIN farmer_pricing_profiles fpp ON fpp.farm_id=f.id AND fpp.crop_year=? "
            "WHERE f.id=?",
            (crop_year, farm_id),
        ).fetchone()
        if not row:
            raise KeyError("Farm not found")
    result = dict(row)
    result.setdefault("volume_discount_pct", 0)
    result.setdefault("early_pay_discount_pct", 0)
    result.setdefault("loyalty_discount_per_unit", 0)
    result.setdefault("custom_discount_per_unit", 0)
    result["crop_year"] = crop_year
    return result


def upsert_farmer_profile(
    farm_id: int,
    crop_year: int,
    volume_discount_pct: float,
    early_pay_discount_pct: float,
    loyalty_discount_per_unit: float,
    custom_discount_per_unit: float,
    pricing_tier: str | None,
    notes: str | None,
    updated_by_user_id: int,
) -> dict[str, Any]:
    with connect() as conn:
        farm = conn.execute("SELECT id FROM farms WHERE id=?", (farm_id,)).fetchone()
        if not farm:
            raise KeyError("Farm not found")
        conn.execute(
            "INSERT INTO farmer_pricing_profiles(farm_id,crop_year,volume_discount_pct,early_pay_discount_pct,"
            "loyalty_discount_per_unit,custom_discount_per_unit,pricing_tier,notes,updated_by_user_id) "
            "VALUES(?,?,?,?,?,?,?,?,?) "
            "ON CONFLICT(farm_id,crop_year) DO UPDATE SET volume_discount_pct=excluded.volume_discount_pct,"
            "early_pay_discount_pct=excluded.early_pay_discount_pct,loyalty_discount_per_unit=excluded.loyalty_discount_per_unit,"
            "custom_discount_per_unit=excluded.custom_discount_per_unit,pricing_tier=excluded.pricing_tier,notes=excluded.notes,"
            "updated_by_user_id=excluded.updated_by_user_id,updated_at=CURRENT_TIMESTAMP",
            (
                farm_id, crop_year, volume_discount_pct, early_pay_discount_pct,
                loyalty_discount_per_unit, custom_discount_per_unit, pricing_tier, notes, updated_by_user_id,
            ),
        )
    return get_farmer_profile(farm_id, crop_year)


def calculated_farmer_price(farm_id: int, crop_year: int, seed_product_id: int) -> dict[str, Any]:
    with connect() as conn:
        row = conn.execute(
            "SELECT f.organization_id,dsp.list_price,dsp.base_price,dsp.dealer_cost,"
            "COALESCE(fpp.volume_discount_pct,0) AS volume_discount_pct,"
            "COALESCE(fpp.early_pay_discount_pct,0) AS early_pay_discount_pct,"
            "COALESCE(fpp.loyalty_discount_per_unit,0) AS loyalty_discount_per_unit,"
            "COALESCE(fpp.custom_discount_per_unit,0) AS custom_discount_per_unit,"
            "fpp.pricing_tier "
            "FROM farms f "
            "LEFT JOIN dealer_seed_prices dsp ON dsp.organization_id=f.organization_id AND dsp.crop_year=? AND dsp.seed_product_id=? AND dsp.status='active' "
            "LEFT JOIN farmer_pricing_profiles fpp ON fpp.farm_id=f.id AND fpp.crop_year=? "
            "WHERE f.id=?",
            (crop_year, seed_product_id, crop_year, farm_id),
        ).fetchone()
        if not row:
            raise KeyError("Farm not found")
    if row.get("base_price") is None:
        return {"available": False, "farm_id": farm_id, "crop_year": crop_year, "seed_product_id": seed_product_id}
    base = float(row["base_price"])
    pct = float(row.get("volume_discount_pct") or 0) + float(row.get("early_pay_discount_pct") or 0)
    per_unit = float(row.get("loyalty_discount_per_unit") or 0) + float(row.get("custom_discount_per_unit") or 0)
    price = max(0.0, base * (1 - pct / 100.0) - per_unit)
    return {
        "available": True,
        "farm_id": farm_id,
        "crop_year": crop_year,
        "seed_product_id": seed_product_id,
        "list_price": _money(row.get("list_price")),
        "base_price": _money(base),
        "dealer_cost": _money(row.get("dealer_cost")),
        "volume_discount_pct": float(row.get("volume_discount_pct") or 0),
        "early_pay_discount_pct": float(row.get("early_pay_discount_pct") or 0),
        "loyalty_discount_per_unit": _money(row.get("loyalty_discount_per_unit") or 0),
        "custom_discount_per_unit": _money(row.get("custom_discount_per_unit") or 0),
        "pricing_tier": row.get("pricing_tier"),
        "calculated_price": _money(price),
    }


def create_price_override(
    field_id: int,
    crop_year: int,
    requested_price: float,
    requested_by_user_id: int,
    request_note: str | None = None,
) -> dict[str, Any]:
    with connect() as conn:
        row = conn.execute(
            "SELECT f.id AS field_id,f.farm_id,fa.organization_id,p.id AS prospect_id,"
            "cp.selected_seed_product_id,cp.units_required "
            "FROM fields f JOIN farms fa ON fa.id=f.farm_id "
            "LEFT JOIN prospects p ON p.farm_id=fa.id "
            "LEFT JOIN field_crop_plans cp ON cp.field_id=f.id AND cp.crop_year=? "
            "WHERE f.id=?",
            (crop_year, field_id),
        ).fetchone()
        if not row:
            raise KeyError("Field not found")
        if not row.get("selected_seed_product_id"):
            raise ValueError("Select a seed product before requesting a price override")
    standard = calculated_farmer_price(int(row["farm_id"]), crop_year, int(row["selected_seed_product_id"]))
    if not standard.get("available"):
        raise ValueError("Dealer pricing is not configured for this product")
    standard_price = float(standard["calculated_price"])
    if abs(float(requested_price) - standard_price) < 0.005:
        raise ValueError("Requested price matches the calculated farmer price; approval is not required")
    with connect() as conn:
        conn.execute(
            "UPDATE price_approval_requests SET status='superseded',updated_at=CURRENT_TIMESTAMP "
            "WHERE field_id=? AND crop_year=? AND status='pending'",
            (field_id, crop_year),
        )
        req = conn.execute(
            "INSERT INTO price_approval_requests(organization_id,farm_id,prospect_id,field_id,crop_year,seed_product_id,"
            "requested_by_user_id,standard_price,requested_price,units,status,request_note) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,'pending',?) RETURNING *",
            (
                row["organization_id"], row["farm_id"], row.get("prospect_id"), field_id, crop_year,
                row["selected_seed_product_id"], requested_by_user_id, standard_price, requested_price,
                row.get("units_required"), request_note,
            ),
        ).fetchone()
    result = dict(req)
    result["standard_pricing"] = standard
    return result


def list_price_requests(organization_id: int, status: str | None = "pending") -> list[dict[str, Any]]:
    where = "WHERE r.organization_id=?"
    params: list[Any] = [organization_id]
    if status:
        where += " AND r.status=?"
        params.append(status)
    with connect() as conn:
        rows = conn.execute(
            "SELECT r.*,f.name AS field_name,fa.farm_name,p.prospect_name,sp.product_name,sp.brand,"
            "u.display_name AS requested_by_name,ru.display_name AS reviewed_by_name,dsp.dealer_cost "
            "FROM price_approval_requests r "
            "JOIN fields f ON f.id=r.field_id JOIN farms fa ON fa.id=r.farm_id "
            "LEFT JOIN prospects p ON p.id=r.prospect_id JOIN seed_products sp ON sp.id=r.seed_product_id "
            "JOIN platform_users u ON u.id=r.requested_by_user_id "
            "LEFT JOIN platform_users ru ON ru.id=r.reviewed_by_user_id "
            "LEFT JOIN dealer_seed_prices dsp ON dsp.organization_id=r.organization_id AND dsp.crop_year=r.crop_year AND dsp.seed_product_id=r.seed_product_id "
            + where + " ORDER BY r.created_at DESC",
            tuple(params),
        ).fetchall()
    return rows_to_dicts(rows)


def review_price_request(
    request_id: int,
    organization_id: int,
    reviewer_user_id: int,
    decision: str,
    approved_price: float | None,
    review_note: str | None,
) -> dict[str, Any]:
    decision = (decision or "").lower()
    if decision not in {"approve", "reject", "counter"}:
        raise ValueError("Decision must be approve, reject, or counter")
    with connect() as conn:
        req = conn.execute(
            "SELECT * FROM price_approval_requests WHERE id=? AND organization_id=?",
            (request_id, organization_id),
        ).fetchone()
        if not req:
            raise KeyError("Price approval request not found")
        if req["status"] != "pending":
            raise ValueError("This price request has already been reviewed")
        final_price = None
        status = "rejected"
        if decision == "approve":
            final_price = float(req["requested_price"]) if approved_price is None else float(approved_price)
            status = "approved"
        elif decision == "counter":
            if approved_price is None:
                raise ValueError("Counter price is required")
            final_price = float(approved_price)
            status = "countered"

        conn.execute(
            "UPDATE price_approval_requests SET status=?,reviewed_by_user_id=?,approved_price=?,review_note=?,"
            "reviewed_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (status, reviewer_user_id, final_price, review_note, request_id),
        )
        if final_price is not None:
            plan = conn.execute(
                "SELECT cp.*,f.acres FROM field_crop_plans cp JOIN fields f ON f.id=cp.field_id "
                "WHERE cp.field_id=? AND cp.crop_year=?",
                (req["field_id"], req["crop_year"]),
            ).fetchone()
            if plan:
                acres = float(plan.get("acres") or 0)
                population = float(plan.get("target_population") or 0)
                seeds_per_unit = float(plan.get("seeds_per_unit") or 0)
                units_required = acres * population / seeds_per_unit if population and seeds_per_unit else None
                seed_cost_per_acre = population / seeds_per_unit * final_price if population and seeds_per_unit else None
                total_seed_cost = units_required * final_price if units_required is not None else None
                conn.execute(
                    "UPDATE field_crop_plans SET seed_price_per_unit=?,pricing_source=?,units_required=?,seed_cost_per_acre=?,"
                    "total_seed_cost=?,updated_at=CURRENT_TIMESTAMP WHERE field_id=? AND crop_year=?",
                    (
                        final_price,
                        "dealer_approved_override" if status == "approved" else "dealer_counter_price",
                        units_required, seed_cost_per_acre, total_seed_cost,
                        req["field_id"], req["crop_year"],
                    ),
                )
        row = conn.execute("SELECT * FROM price_approval_requests WHERE id=?", (request_id,)).fetchone()
    return row_to_dict(row) or {}


def latest_field_price_request(field_id: int, crop_year: int) -> dict[str, Any] | None:
    with connect() as conn:
        row = conn.execute(
            "SELECT r.*,u.display_name AS requested_by_name,ru.display_name AS reviewed_by_name "
            "FROM price_approval_requests r JOIN platform_users u ON u.id=r.requested_by_user_id "
            "LEFT JOIN platform_users ru ON ru.id=r.reviewed_by_user_id "
            "WHERE r.field_id=? AND r.crop_year=? ORDER BY r.created_at DESC LIMIT 1",
            (field_id, crop_year),
        ).fetchone()
    return row_to_dict(row)
